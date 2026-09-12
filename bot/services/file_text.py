from __future__ import annotations

import io
import re
from typing import Optional

from telegram import Bot

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def clean_text(text: str) -> str:
    """
    تنظيف النص المستخرج من الملفات.
    """

    if not text:
        return ""

    text = text.replace("\x00", " ")

    # توحيد الأسطر
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # إزالة المسافات الزائدة
    text = re.sub(r"[ \t]+", " ", text)

    # إزالة الأسطر الفارغة المتكررة
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def extract_pdf(data: bytes) -> str:
    """
    استخراج النص من PDF.
    """

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("تعذر فتح ملف PDF.") from exc

    pages = []

    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        text = clean_text(text)

        if text:
            pages.append(
                f"--- الصفحة {index} ---\n{text}"
            )

    return clean_text("\n\n".join(pages))


def extract_docx(data: bytes) -> str:
    """
    استخراج النص من Word DOCX.
    """

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("تعذر فتح ملف Word.") from exc

    parts = []

    # الفقرات
    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)

        if text:
            parts.append(text)

    # الجداول
    for table in document.tables:
        rows = []

        for row in table.rows:
            cells = []

            for cell in row.cells:
                text = clean_text(cell.text)

                if text:
                    cells.append(text)

            if cells:
                rows.append(" | ".join(cells))

        if rows:
            parts.append(
                "--- جدول ---\n"
                + "\n".join(rows)
            )

    return clean_text("\n\n".join(parts))


def extract_pptx(data: bytes) -> str:
    """
    استخراج النص من PowerPoint.
    """

    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("تعذر فتح ملف PowerPoint.") from exc

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_parts = []

        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue

            try:
                text = clean_text(shape.text)
            except Exception:
                text = ""

            if text:
                slide_parts.append(text)

        if slide_parts:
            slides.append(
                f"--- الشريحة {slide_number} ---\n"
                + "\n".join(slide_parts)
            )

    return clean_text("\n\n".join(slides))


def extract_plain_text(data: bytes) -> str:
    """
    استخراج النص من الملفات النصية العادية.
    """

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1256",
        "latin-1",
    ]

    for encoding in encodings:
        try:
            text = data.decode(encoding)

            if text.strip():
                return clean_text(text)

        except UnicodeDecodeError:
            continue

    raise ValueError(
        "تعذر قراءة محتوى الملف النصي."
    )


def extract_text_from_bytes(
    data: bytes,
    file_type: Optional[str] = None,
    file_name: Optional[str] = None,
) -> str:
    """
    تحديد نوع الملف واستخراج النص منه.
    """

    normalized_type = (
        (file_type or "")
        .strip()
        .lower()
    )

    normalized_name = (
        (file_name or "")
        .strip()
        .lower()
    )

    if (
        normalized_type == "pdf"
        or normalized_name.endswith(".pdf")
        or normalized_name.endswith(".PDF".lower())
    ):
        return extract_pdf(data)

    if (
        normalized_type in {
            "docx",
            "word",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        }
        or normalized_name.endswith(".docx")
    ):
        return extract_docx(data)

    if (
        normalized_type in {
            "pptx",
            "powerpoint",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        }
        or normalized_name.endswith(".pptx")
    ):
        return extract_pptx(data)

    if normalized_type in {
        "txt",
        "text",
        "text/plain",
    } or normalized_name.endswith(".txt"):
        return extract_plain_text(data)

    # محاولة تلقائية حسب محتوى الملف
    if data.startswith(b"%PDF"):
        return extract_pdf(data)

    # DOCX و PPTX عبارة عن ZIP
    if data.startswith(b"PK"):
        try:
            return extract_docx(data)
        except Exception:
            pass

        try:
            return extract_pptx(data)
        except Exception:
            pass

    return extract_plain_text(data)


async def download_telegram_file(
    bot: Bot,
    telegram_file_id: str,
) -> bytes:
    """
    تحميل ملف من Telegram باستخدام file_id.
    """

    if not telegram_file_id:
        raise ValueError(
            "معرّف ملف Telegram غير موجود."
        )

    try:
        telegram_file = await bot.get_file(
            telegram_file_id
        )
    except Exception as exc:
        raise ValueError(
            "تعذر الوصول إلى الملف الموجود في Telegram."
        ) from exc

    buffer = io.BytesIO()

    try:
        await telegram_file.download_to_memory(
            buffer
        )
    except Exception as exc:
        raise ValueError(
            "تعذر تحميل الملف من Telegram."
        ) from exc

    return buffer.getvalue()
